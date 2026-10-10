"""Build a polished .docx matching the Nxtwave Edge 2.0 sample feedback form."""

from __future__ import annotations

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from config import ASSESSMENT_META, Cutoffs
from feedback_builder import _display_rows, _fmt_score, _overall_status
from processor import StudentData

# Sample-inspired palette
NAVY = RGBColor(0x1B, 0x3A, 0x5F)
HEADER_BG = "1B3A5F"
ALT_ROW_BG = "F3F6FA"
RESULT_BG = "FFF3CD"
RESULT_BORDER = "E6A800"
FOOTER_BG = "E8F1FB"
WHITE = "FFFFFF"


def _set_run_font(run, *, size=11, bold=False, color=None, name="Calibri"):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    if color is not None:
        run.font.color.rgb = color
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)


def _shade_cell(cell, hex_color: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def _set_cell_border(cell, color: str = "CCCCCC", sz: str = "8") -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), sz)
        element.set(qn("w:color"), color)
        tcBorders.append(element)
    tcPr.append(tcBorders)


def _clear_cell(cell) -> None:
    for p in cell.paragraphs:
        p.clear()


def _cell_text(cell, text: str, *, bold=False, size=10, color=None, center=False) -> None:
    _clear_cell(cell)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    _set_run_font(run, size=size, bold=bold, color=color)


def _add_heading(doc: Document, text: str, *, size=16, color=NAVY) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    _set_run_font(run, size=size, bold=True, color=color)
    p.paragraph_format.space_after = Pt(4)


def _add_subheading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    _set_run_font(run, size=13, bold=True, color=NAVY)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)


def _add_body(doc: Document, text: str, *, size=10, italic=False, center=False) -> None:
    p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    _set_run_font(run, size=size, color=RGBColor(0x33, 0x33, 0x33))
    run.italic = italic
    p.paragraph_format.space_after = Pt(4)


def _add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    _set_run_font(run, size=10, color=RGBColor(0x33, 0x33, 0x33))
    p.paragraph_format.space_after = Pt(3)


def _info_table(doc: Document, student: StudentData) -> None:
    table = doc.add_table(rows=2, cols=2)
    table.autofit = True
    rows = [
        ("Learner", student.user_name),
        ("Assessment date", student.assessment_date or "Not available"),
    ]
    for i, (label, value) in enumerate(rows):
        _cell_text(table.rows[i].cells[0], label, bold=True, size=10, color=NAVY)
        _cell_text(table.rows[i].cells[1], value, size=10)
        _shade_cell(table.rows[i].cells[0], ALT_ROW_BG)
        for cell in table.rows[i].cells:
            _set_cell_border(cell, "D0D7DE", "6")
    doc.add_paragraph()


def _result_box(doc: Document, result_line: str, score_line: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    _shade_cell(cell, RESULT_BG)
    _set_cell_border(cell, RESULT_BORDER, "16")
    _clear_cell(cell)
    p1 = cell.paragraphs[0]
    r1 = p1.add_run(result_line)
    _set_run_font(r1, size=11, bold=True, color=RGBColor(0x5C, 0x3D, 0x00))
    p2 = cell.add_paragraph()
    r2 = p2.add_run(score_line)
    _set_run_font(r2, size=11, bold=True, color=RGBColor(0x5C, 0x3D, 0x00))
    doc.add_paragraph()


def _scores_table(doc: Document, display: list[dict], overall, overall_max, status_line: str) -> None:
    table = doc.add_table(rows=1 + len(display) + 1, cols=4)
    headers = ["Section", "Your score", "Out of", "Status"]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        _cell_text(cell, h, bold=True, size=10, color=RGBColor(0xFF, 0xFF, 0xFF), center=(i > 0))
        _shade_cell(cell, HEADER_BG)
        _set_cell_border(cell, HEADER_BG, "8")

    for r, row in enumerate(display, start=1):
        values = [
            f"{r}. {row['label']}",
            _fmt_score(row["score"]),
            str(row["max_marks"]),
            row["status"],
        ]
        for c, val in enumerate(values):
            cell = table.rows[r].cells[c]
            _cell_text(cell, val, size=10, center=(c > 0), bold=(c == 3))
            if r % 2 == 0:
                _shade_cell(cell, ALT_ROW_BG)
            _set_cell_border(cell, "D0D7DE", "6")

    total_row = table.rows[-1]
    totals = ["Total", _fmt_score(overall), str(overall_max), status_line]
    for c, val in enumerate(totals):
        cell = total_row.cells[c]
        _cell_text(cell, val, bold=True, size=10, center=(c > 0), color=NAVY)
        _shade_cell(cell, "E8EEF5")
        _set_cell_border(cell, "D0D7DE", "6")

    note = doc.add_paragraph()
    run = note.add_run("Status is Cleared or Below the bar for each section.")
    _set_run_font(run, size=9, color=RGBColor(0x66, 0x66, 0x66))
    run.italic = True


def _topics_table(doc: Document, topic_rows: list[dict]) -> None:
    if not topic_rows:
        _add_body(doc, "Topic-level breakdown unavailable for this attempt.", italic=True)
        return
    table = doc.add_table(rows=1 + len(topic_rows), cols=3)
    for i, h in enumerate(["Section", "Topic area", "Your level"]):
        cell = table.rows[0].cells[i]
        _cell_text(cell, h, bold=True, size=10, color=RGBColor(0xFF, 0xFF, 0xFF))
        _shade_cell(cell, HEADER_BG)
        _set_cell_border(cell, HEADER_BG, "8")

    level_colors = {
        "strong": RGBColor(0x1B, 0x7A, 0x3A),
        "partial": RGBColor(0xB3, 0x6B, 0x00),
        "needs work": RGBColor(0xB0, 0x2A, 0x37),
        "not attempted": RGBColor(0x66, 0x66, 0x66),
    }
    for r, row in enumerate(topic_rows, start=1):
        level = str(row.get("level") or "")
        color = level_colors.get(level.strip().lower(), RGBColor(0x33, 0x33, 0x33))
        vals = [row.get("section", ""), row.get("topic_area", ""), level]
        for c, val in enumerate(vals):
            cell = table.rows[r].cells[c]
            _cell_text(
                cell,
                str(val),
                size=10,
                bold=(c == 2),
                color=color if c == 2 else None,
                center=(c == 2),
            )
            if r % 2 == 0:
                _shade_cell(cell, ALT_ROW_BG)
            _set_cell_border(cell, "D0D7DE", "6")


def _footer_box(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    _shade_cell(cell, FOOTER_BG)
    _set_cell_border(cell, "7BA7D9", "12")
    _clear_cell(cell)
    p = cell.paragraphs[0]
    run = p.add_run("About this report: Scores are generated from your submitted answers and are final.")
    _set_run_font(run, size=9, color=RGBColor(0x1B, 0x3A, 0x5F))


def build_feedback_docx_bytes(
    student: StudentData,
    analysis: dict,
    cutoffs: Cutoffs,
) -> bytes:
    scores = analysis.get("_scores") or {}
    overall = scores.get("overall_score")
    overall_max = scores.get("overall_max", 190)
    display = _display_rows(scores, cutoffs)
    status_line = _overall_status(overall, cutoffs)

    if analysis.get("result_line"):
        result_line = analysis["result_line"].strip()
        if not result_line.lower().startswith("result"):
            result_line = f"Result — {result_line}"
    else:
        result_line = f"Result — {status_line}"

    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(1.6)
        section.bottom_margin = Cm(1.6)
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.8)

    _add_heading(doc, ASSESSMENT_META["title"], size=14)
    _add_heading(doc, "Your Assessment Report", size=18)
    _add_body(doc, ASSESSMENT_META["duration_label"], size=10, italic=True, center=True)

    _info_table(doc, student)
    _result_box(
        doc,
        result_line,
        f"You scored {_fmt_score(overall)} out of {overall_max}.",
    )

    _add_subheading(doc, "Your scores")
    _scores_table(doc, display, overall, overall_max, status_line)

    _add_subheading(doc, "Where you stood, topic by topic")
    _add_body(
        doc,
        "Each topic below is marked Strong, Partial, Needs work or Not Attempted, "
        "based on the marks you scored in that part of the paper.",
        size=9,
        italic=True,
    )
    _topics_table(doc, analysis.get("topic_rows") or [])

    _add_subheading(doc, "Highlights from this attempt")
    for bullet in analysis.get("what_went_well") or [
        "No clear topic strengths could be evidenced from available data."
    ]:
        _add_bullet(doc, bullet)

    _add_subheading(doc, "Where the score slipped")
    for bullet in analysis.get("where_you_lost_marks") or [
        "Key areas to improve could not be determined from available data."
    ]:
        _add_bullet(doc, bullet)

    _add_subheading(doc, "Focus plan for your next attempt")
    for bullet in analysis.get("what_to_work_on_next") or [
        "Focus on the weaker topic areas listed above."
    ]:
        _add_bullet(doc, bullet)

    doc.add_paragraph()
    _footer_box(doc)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
